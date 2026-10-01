import openpyxl

wb = openpyxl.load_workbook('c:/svadd/Billing_Format.xlsx', data_only=False)
ws = wb['Biling Sheet']

print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")

for r in range(100, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    if any(v is not None for v in vals):
        print(f"Row {r:3d}: {vals}")

# Also check data validation / dropdowns!
print("\n--- Data Validations ---")
for dv in ws.data_validations.dataValidation:
    print("Formula1:", dv.formula1, "Formula2:", dv.formula2, "Sqref:", dv.sqref, "Type:", dv.type)

# Also check defined names / named ranges
print("\n--- Defined Names ---")
for name in wb.defined_names.definedName:
    print(name.name, name.value)
